# godomall_register/market_category_fetch.py
"""OC 양식 Excel → oc_category_markets 테이블 수집."""
from __future__ import annotations

import sqlite3
from pathlib import Path


def _parse_market_cat_string(s: str) -> dict[str, dict]:
    """'market,code,name\n...' → {market: {code, name}}."""
    result = {}
    for line in (s or "").strip().splitlines():
        parts = line.strip().split(",", 2)
        if len(parts) >= 2:
            market, code = parts[0].strip(), parts[1].strip()
            name = parts[2].strip() if len(parts) > 2 else ""
            if market and code:
                result[market] = {"code": code, "name": name}
    return result


def fetch_from_oc_excel(excel_path: str, catalog_db_path: str) -> dict:
    """OC 양식 Excel(OWNERCLAN 시트)에서 카테고리+마켓카테고리 추출 → oc_category_markets 저장.

    Returns: {"categories": int, "rows_inserted": int, "errors": list}
    """
    import openpyxl
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    if "OWNERCLAN" not in wb.sheetnames:
        raise ValueError(f"'OWNERCLAN' 시트 없음. 시트목록: {wb.sheetnames}")
    ws = wb["OWNERCLAN"]

    rows = list(ws.iter_rows(values_only=True))
    # Row 0: 섹션 헤더 (상품기본정보, 카테고리정보, ...)
    # Row 1: 컬럼 헤더 (판매자상품코드1, ..., 카테고리코드, 카테고리명, 마켓카테고리, ...)
    if len(rows) < 3:
        return {"categories": 0, "rows_inserted": 0, "errors": ["데이터 행 없음"]}

    headers = list(rows[1])
    col = {str(h): i for i, h in enumerate(headers) if h}
    needed = {"카테고리코드", "카테고리명", "마켓카테고리"}
    missing = needed - col.keys()
    if missing:
        raise ValueError(f"필수 컬럼 없음: {missing}. 헤더: {headers[:10]}")

    ki = col["카테고리코드"]
    ni = col["카테고리명"]
    mi = col["마켓카테고리"]

    # 카테고리 단위 수집 (상품별 중복 제거)
    cats: dict[str, dict] = {}
    for row in rows[2:]:
        key = str(row[ki] or "").strip()
        name = str(row[ni] or "").strip()
        mcat = str(row[mi] or "").strip()
        if key and key not in cats:
            cats[key] = {"name": name, "mcat": mcat}

    # oc_category_markets 저장
    con = sqlite3.connect(catalog_db_path)
    con.execute("PRAGMA journal_mode=WAL")
    con.execute("""
        CREATE TABLE IF NOT EXISTS oc_category_markets (
            oc_category_key  TEXT NOT NULL,
            oc_category_name TEXT,
            market           TEXT NOT NULL,
            market_cat_code  TEXT,
            market_cat_name  TEXT,
            is_manual        INTEGER DEFAULT 0,
            updated_at       TEXT DEFAULT (datetime('now')),
            PRIMARY KEY (oc_category_key, market)
        )
    """)

    rows_inserted = 0
    errors = []
    for cat_key, info in cats.items():
        markets = _parse_market_cat_string(info["mcat"])
        if not markets:
            errors.append(f"{cat_key}: 마켓카테고리 없음")
            continue
        for market, minfo in markets.items():
            con.execute("""
                INSERT OR REPLACE INTO oc_category_markets
                (oc_category_key, oc_category_name, market,
                 market_cat_code, market_cat_name, updated_at)
                VALUES (?, ?, ?, ?, ?, datetime('now'))
            """, [cat_key, info["name"], market, minfo["code"], minfo["name"]])
            rows_inserted += 1

    con.commit()
    con.close()
    return {"categories": len(cats), "rows_inserted": rows_inserted, "errors": errors}
