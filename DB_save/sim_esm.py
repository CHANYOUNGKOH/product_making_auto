"""ESM (옥션/지마켓) 3-전략 시뮬레이션 엑셀 생성.

A타입 (유배, J=2%)  : 최저가전략 + 일반판매전략 + 광고전략
B타입 (무배, J=61%) :            일반판매전략 + 광고전략

실행:
    cd .worktrees/feat-pricing
    python DB_save/sim_esm.py
"""
from __future__ import annotations

import sys
import os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from DB_save.price_engine import solve_h, solve_h_by_v, _resolve_target_s
from DB_save.pricing_strategies import get_pricing_strategy

# ─── 정책 (A타입 유배 / B타입 무배) ─────────────────────────────────────────
POLICY_A = {                        # 유배, J=2%
    "commission_rate": 16.5,
    "commission_base": "pre_discount",
    "discount_rate": 2.0,
    "coupon_amount": 0,
    "reward_rate": 0,
}
POLICY_B = {                        # 무배, J=61%
    "commission_rate": 16.5,
    "commission_base": "pre_discount",
    "discount_rate": 61.0,
    "coupon_amount": 0,
    "reward_rate": 0,
}
SHIPPING_A = 0       # 유배: 구매자 별도 부담
SHIPPING_B = 3000    # 무배: 판매자 흡수

# ─── 샘플 상품 (28개) ────────────────────────────────────────────────────────
COSTS = [
    50, 100, 150, 200, 300, 400, 500,
    700, 1000, 1500, 2000, 3000, 4000, 5000,
    7000, 8000, 10000, 15000, 20000, 30000,
    50000, 70000, 100000, 150000, 200000, 300000, 400000, 500000,
]
PRODUCTS = [
    {"code": f"TEST{i+1:03d}", "cost_a": float(c)}
    for i, c in enumerate(COSTS)
]

# ─── 스타일 ─────────────────────────────────────────────────────────────────
HDR_FONT   = Font(bold=True, size=10, color="FFFFFF")
HDR_ORANGE = PatternFill(start_color="C55A11", end_color="C55A11", fill_type="solid")  # 최저가
HDR_BLUE   = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")  # 일반A
HDR_LBLUE  = PatternFill(start_color="9DC3E6", end_color="9DC3E6", fill_type="solid")  # 일반B
HDR_GREEN  = PatternFill(start_color="375623", end_color="375623", fill_type="solid")  # 광고A
HDR_LGREEN = PatternFill(start_color="A9D18E", end_color="A9D18E", fill_type="solid")  # 광고B
CLAMP_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
THIN = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

# 컬럼 정의
BASE_COLS = [
    ("상품코드",       14),
    ("A 원가",        10),
    ("C 마진배수",     10),
    ("J 할인율",        8),
    ("배송비흡수",     10),
    ("G",             12),
    ("H 판매가율",     10),
    ("I 마켓등록가",   14),
    ("L 할인후가격",   14),
    ("N 마켓수수료",   14),
    ("O 입금액",       12),
    ("S 영업이익",     12),
    ("target",         10),
    ("V 매출대비%",    10),
    ("마켓등록가(10원)", 16),
    ("상태",           12),
]

FMT_INT = "#,##0"
FMT_2DP = "0.00"
FMT_PCT = "0.0%"


def _setup_headers(ws, fill: PatternFill) -> None:
    for col_idx, (name, width) in enumerate(BASE_COLS, 1):
        c = ws.cell(row=1, column=col_idx, value=name)
        c.font = HDR_FONT
        c.fill = fill
        c.border = THIN
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 35


def _write_row(ws, row_idx: int, row: list, status: str) -> None:
    for col_idx, val in enumerate(row, 1):
        c = ws.cell(row=row_idx, column=col_idx, value=val)
        c.border = THIN
        if col_idx in (2, 5, 6, 8, 9, 10, 11, 12, 15):
            c.number_format = FMT_INT
        elif col_idx in (3, 7):
            c.number_format = FMT_2DP
        elif col_idx in (4, 13, 14):
            c.number_format = FMT_PCT
        if status not in ("ok",) and col_idx in (7, 16):
            c.fill = CLAMP_FILL


def _fill_lowest(ws, strategy: dict, policy: dict, shipping: float) -> None:
    """최저가전략: absolute_s, A타입(J=2%)만."""
    margin_c  = strategy["margin_c"]
    min_h     = strategy["min_h"]
    max_h     = strategy["max_h"]
    round_h   = strategy["round_h"]
    round_mode = strategy.get("round_mode", "nearest")
    bands     = strategy["bands"]

    for row_idx, prod in enumerate(PRODUCTS, 2):
        cost_a = prod["cost_a"]
        target_s = _resolve_target_s(cost_a, bands)
        sol = solve_h(
            cost_a=cost_a, margin_c=margin_c, policy=policy,
            target_s=target_s, shipping_absorbed=shipping,
            min_h=min_h, max_h=max_h, round_h=round_h, round_mode=round_mode,
        )
        h   = sol["h"] or 0
        fwd = sol["forward"] or {}
        V_pct = fwd.get("revenue_ratio", 0) or 0

        row = [
            prod["code"], cost_a, margin_c,
            policy["discount_rate"] / 100,
            shipping,
            fwd.get("G", 0), h,
            fwd.get("I", 0), fwd.get("L", 0),
            fwd.get("N", 0), fwd.get("O", 0), fwd.get("S", 0),
            target_s, V_pct / 100,
            fwd.get("market_price", 0), sol["status"],
        ]
        _write_row(ws, row_idx, row, sol["status"])


def _fill_v(ws, strategy: dict, policy: dict, shipping: float) -> None:
    """일반판매/광고전략: v_percent."""
    margin_c = strategy["margin_c"]
    min_h    = strategy["min_h"]
    max_h    = strategy["max_h"]
    round_h  = strategy["round_h"]
    bands    = strategy["bands"]

    for row_idx, prod in enumerate(PRODUCTS, 2):
        cost_a   = prod["cost_a"]
        target_v = _resolve_target_s(cost_a, bands)
        sol = solve_h_by_v(
            cost_a=cost_a, margin_c=margin_c, policy=policy,
            target_v_pct=target_v, shipping_absorbed=shipping,
            min_h=min_h, max_h=max_h, round_h=round_h,
        )
        h   = sol["h"] or 0
        fwd = sol["forward"] or {}
        V_pct = fwd.get("revenue_ratio", 0) or 0

        row = [
            prod["code"], cost_a, margin_c,
            policy["discount_rate"] / 100,
            shipping,
            fwd.get("G", 0), h,
            fwd.get("I", 0), fwd.get("L", 0),
            fwd.get("N", 0), fwd.get("O", 0), fwd.get("S", 0),
            target_v / 100, V_pct / 100,
            fwd.get("market_price", 0), sol["status"],
        ]
        _write_row(ws, row_idx, row, sol["status"])


def _add_policy_sheet(wb, market: str, strategies: list) -> None:
    ws = wb.create_sheet("정책정보")
    rows = [
        ("마켓", market),
        ("수수료율", "16.5%"),
        ("수수료기준", "pre_discount (등록가 기준)"),
        ("", ""),
        ("[A타입] 할인율 J", "2%"),
        ("[A타입] 배송비", "유배 (구매자 3,000원 별도)"),
        ("[B타입] 할인율 J", "61%"),
        ("[B타입] 배송비", "무배 (판매자 3,000원 흡수)"),
        ("", ""),
    ]
    for sid, strat, policy, shipping in strategies:
        rows.append((f"[{sid}] C", strat["margin_c"]))
        rows.append((f"[{sid}] J", f"{policy['discount_rate']}%"))
        rows.append((f"[{sid}] shipping", f"{shipping}원"))
        rows.append((f"[{sid}] bands", str(strat["bands"])))
        rows.append(("", ""))

    for r, (k, v) in enumerate(rows, 1):
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=str(v))
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100


def main(market: str = "옥션") -> None:
    strat_low  = get_pricing_strategy(market, "lowest_price")
    strat_norm = get_pricing_strategy(market, "normal_sale")
    strat_ad   = get_pricing_strategy(market, "cpc_ad")

    wb = openpyxl.Workbook()

    # Sheet 1: 최저가전략 (A타입, J=2%)
    ws1 = wb.active
    ws1.title = "최저가(유배A)"
    _setup_headers(ws1, HDR_ORANGE)
    _fill_lowest(ws1, strat_low, POLICY_A, SHIPPING_A)

    # Sheet 2: 일반판매전략 A타입
    ws2 = wb.create_sheet("일반판매(유배A)")
    _setup_headers(ws2, HDR_BLUE)
    _fill_v(ws2, strat_norm, POLICY_A, SHIPPING_A)

    # Sheet 3: 일반판매전략 B타입
    ws3 = wb.create_sheet("일반판매(무배B)")
    _setup_headers(ws3, HDR_LBLUE)
    _fill_v(ws3, strat_norm, POLICY_B, SHIPPING_B)

    # Sheet 4: 광고전략 A타입
    ws4 = wb.create_sheet("광고(유배A)")
    _setup_headers(ws4, HDR_GREEN)
    _fill_v(ws4, strat_ad, POLICY_A, SHIPPING_A)

    # Sheet 5: 광고전략 B타입
    ws5 = wb.create_sheet("광고(무배B)")
    _setup_headers(ws5, HDR_LGREEN)
    _fill_v(ws5, strat_ad, POLICY_B, SHIPPING_B)

    _add_policy_sheet(wb, market, [
        ("최저가(유배A)",   strat_low,  POLICY_A, SHIPPING_A),
        ("일반판매(유배A)", strat_norm, POLICY_A, SHIPPING_A),
        ("일반판매(무배B)", strat_norm, POLICY_B, SHIPPING_B),
        ("광고(유배A)",     strat_ad,   POLICY_A, SHIPPING_A),
        ("광고(무배B)",     strat_ad,   POLICY_B, SHIPPING_B),
    ])

    out = f"DB_save/{market}_3전략_시뮬_v1.xlsx"
    wb.save(out)
    print(f"Saved: {out}")

    # ── 콘솔 요약 (최저가전략 A타입) ─────────────────────────────────────────
    print(f"\n=== {market} 최저가전략 (J=2%, 유배, absolute_s) ===")
    print(f"{'code':<10}{'cost':>8}{'target_S':>10}{'H_raw':>8}{'H':>6}"
          f"{'I':>10}{'L':>10}{'S':>9}{'V%':>7}  status")
    print("-" * 88)
    for prod in PRODUCTS:
        cost_a   = prod["cost_a"]
        target_s = _resolve_target_s(cost_a, strat_low["bands"])
        sol = solve_h(
            cost_a=cost_a, margin_c=strat_low["margin_c"], policy=POLICY_A,
            target_s=target_s, shipping_absorbed=SHIPPING_A,
            min_h=strat_low["min_h"], max_h=strat_low["max_h"],
            round_h=strat_low["round_h"],
            round_mode=strat_low.get("round_mode", "nearest"),
        )
        fwd = sol["forward"] or {}
        print(f"{prod['code']:<10}{cost_a:>8,.0f}{target_s:>10,.0f}"
              f"{sol['h_raw'] or 0:>8.4f}{sol['h'] or 0:>6.2f}"
              f"{fwd.get('I',0):>10,.0f}{fwd.get('L',0):>10,.0f}"
              f"{fwd.get('S',0):>9,.0f}{fwd.get('revenue_ratio',0):>6.1f}%"
              f"  {sol['status']}")


if __name__ == "__main__":
    main("옥션")
    main("지마켓")
