"""스마트스토어 3-전략 시뮬레이션 엑셀 생성.

실행:
    cd .worktrees/feat-pricing
    python DB_save/sim_smartstore.py
"""
from __future__ import annotations

import sys
import os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import math
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from DB_save.price_engine import forward_calc, solve_h, solve_h_by_v, _resolve_target_s
from DB_save.pricing_strategies import get_pricing_strategy

# ─── 스마트스토어 정책 ──────────────────────────────────────────────────────
POLICY = {
    "commission_rate": 6.5,
    "commission_base": "post_discount",
    "discount_rate": 50,       # sales_channels.json 의 default_discount_rate
    "coupon_amount": 150,
    "reward_rate": 15,
}

# ─── 샘플 상품 (28개, 원가 50~500,000원) ────────────────────────────────────
COSTS = [
    50, 100, 150, 200, 300, 400, 500,
    700, 1000, 1500, 2000, 3000, 4000, 5000,
    7000, 8000, 10000, 15000, 20000, 30000,
    50000, 70000, 100000, 150000, 200000, 300000, 400000, 500000,
]

PRODUCTS = [
    {"code": f"TEST{i+1:03d}", "cost_a": float(c), "shipping": 0.0}
    for i, c in enumerate(COSTS)
]

# ─── 스타일 ─────────────────────────────────────────────────────────────────
HDR_FONT = Font(bold=True, size=10, color="000000")
HDR_ORANGE = PatternFill(start_color="F4B942", end_color="F4B942", fill_type="solid")   # 최저가
HDR_BLUE   = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")   # 일반판매
HDR_GREEN  = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")   # 광고
CLAMP_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
THIN = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

# ─── 컬럼 정의 (최저가용, 일반/광고용 공통) ─────────────────────────────────
BASE_COLS = [
    ("상품코드",          14),
    ("A 원가",           10),
    ("C 마진배수",        10),
    ("E 쿠폰",            8),
    ("G 원가×C",         12),   # display: cost*C (E 제외)
    ("H 판매가율",        10),
    ("I 마켓등록가",      14),
    ("J 할인율",          10),
    ("L 할인후가격",      14),
    ("M 쿠폰정산",        10),
    ("N 마켓수수료",      14),
    ("P 스마일캐시%",     12),
    ("Q 스마일캐시액",    12),
    ("O 입금액",          12),
    ("S 영업이익",        12),
    ("target",            10),   # target_S or target_V%
    ("V 매출대비%",       10),
    ("마켓등록가(10원)",   16),
    ("상태",              12),
]

FMT_INT = "#,##0"
FMT_2DP = "0.00"
FMT_PCT = "0.0%"
FMT_4DP = "0.0000"


def _hdr_fill(strategy_id: str) -> PatternFill:
    return {
        "lowest_price": HDR_ORANGE,
        "normal_sale":  HDR_BLUE,
        "cpc_ad":       HDR_GREEN,
    }[strategy_id]


def _setup_headers(ws, strategy_id: str) -> None:
    fill = _hdr_fill(strategy_id)
    for col_idx, (name, width) in enumerate(BASE_COLS, 1):
        c = ws.cell(row=1, column=col_idx, value=name)
        c.font = HDR_FONT
        c.fill = fill
        c.border = THIN
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 35


def _fill_lowest_price(ws, strategy: dict) -> None:
    """최저가전략: absolute_s metric → solve_h."""
    margin_c = strategy["margin_c"]
    min_h    = strategy["min_h"]
    max_h    = strategy["max_h"]
    round_h  = strategy["round_h"]
    bands    = strategy["bands"]

    J = POLICY["discount_rate"] / 100
    E = POLICY["coupon_amount"]
    P = POLICY["reward_rate"] / 100

    for row_idx, prod in enumerate(PRODUCTS, 2):
        cost_a = prod["cost_a"]

        target_s = _resolve_target_s(cost_a, bands)
        sol = solve_h(
            cost_a=cost_a, margin_c=margin_c, policy=POLICY,
            target_s=target_s, shipping_absorbed=0,
            min_h=min_h, max_h=max_h, round_h=round_h,
        )
        h  = sol["h"] or 0
        fwd = sol["forward"] or {}

        G_display = cost_a * margin_c           # G 표시: 원가×C (E 제외)
        V_pct = fwd.get("revenue_ratio", 0) or 0

        row = [
            prod["code"],                        # 상품코드
            cost_a,                              # A 원가
            margin_c,                            # C 마진배수
            E,                                   # E 쿠폰
            G_display,                           # G 원가×C (표시용)
            h,                                   # H 판매가율
            fwd.get("I", 0),                     # I 마켓등록가
            J,                                   # J 할인율
            fwd.get("L", 0),                     # L 할인후가격
            E,                                   # M 쿠폰정산 (=E)
            fwd.get("N", 0),                     # N 마켓수수료
            P,                                   # P 스마일캐시%
            fwd.get("Q", 0),                     # Q 스마일캐시액
            fwd.get("O", 0),                     # O 입금액
            fwd.get("S", 0),                     # S 영업이익
            target_s,                            # target_S
            V_pct / 100,                         # V 매출대비%
            fwd.get("market_price", 0),          # 마켓등록가(10원)
            sol["status"],                       # 상태
        ]
        _write_row(ws, row_idx, row, strategy_id="lowest_price", status=sol["status"])


def _fill_v_strategy(ws, strategy: dict) -> None:
    """일반판매/광고전략: v_percent metric → solve_h_by_v."""
    margin_c = strategy["margin_c"]
    min_h    = strategy["min_h"]
    max_h    = strategy["max_h"]
    round_h  = strategy["round_h"]
    bands    = strategy["bands"]

    J = POLICY["discount_rate"] / 100
    E = POLICY["coupon_amount"]
    P = POLICY["reward_rate"] / 100

    for row_idx, prod in enumerate(PRODUCTS, 2):
        cost_a = prod["cost_a"]

        target_v = _resolve_target_s(cost_a, bands)   # reuse helper (v_percent band)
        sol = solve_h_by_v(
            cost_a=cost_a, margin_c=margin_c, policy=POLICY,
            target_v_pct=target_v, shipping_absorbed=0,
            min_h=min_h, max_h=max_h, round_h=round_h,
        )
        h  = sol["h"] or 0
        fwd = sol["forward"] or {}

        G_display = cost_a * margin_c
        V_pct = fwd.get("revenue_ratio", 0) or 0

        row = [
            prod["code"],
            cost_a,
            margin_c,
            E,
            G_display,
            h,
            fwd.get("I", 0),
            J,
            fwd.get("L", 0),
            E,
            fwd.get("N", 0),
            P,
            fwd.get("Q", 0),
            fwd.get("O", 0),
            fwd.get("S", 0),
            target_v / 100,                       # target_V% (소수)
            V_pct / 100,
            fwd.get("market_price", 0),
            sol["status"],
        ]
        _write_row(ws, row_idx, row, strategy_id="v_strategy", status=sol["status"])


def _write_row(ws, row_idx: int, row: list, strategy_id: str, status: str) -> None:
    for col_idx, val in enumerate(row, 1):
        c = ws.cell(row=row_idx, column=col_idx, value=val)
        c.border = THIN

        # 숫자 포맷
        if col_idx in (2, 5, 7, 9, 10, 11, 13, 14, 15, 18):
            c.number_format = FMT_INT
        elif col_idx in (3, 6):
            c.number_format = FMT_2DP
        elif col_idx in (8, 12, 16, 17):
            c.number_format = FMT_PCT
        elif col_idx == 4:
            c.number_format = FMT_INT

        # 경고 색상
        if status != "ok" and col_idx in (6, 19):
            c.fill = CLAMP_FILL


def _add_policy_sheet(wb: openpyxl.Workbook, strategies: list[tuple[str, dict]]) -> None:
    ws = wb.create_sheet("정책정보")
    rows = [
        ("마켓", "스마트스토어"),
        ("수수료율", f"{POLICY['commission_rate']}%"),
        ("수수료기준", POLICY["commission_base"]),
        ("할인율 J", f"{POLICY['discount_rate']}%"),
        ("쿠폰 E", f"{POLICY['coupon_amount']}원"),
        ("스마일캐시 P", f"{POLICY['reward_rate']}%"),
        ("", ""),
    ]
    for sid, strat in strategies:
        rows.append((f"[{sid}] C", strat["margin_c"]))
        rows.append((f"[{sid}] metric", strat["metric"]))
        rows.append((f"[{sid}] bands", str(strat["bands"])))
        rows.append(("", ""))

    for r, (k, v) in enumerate(rows, 1):
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=str(v))
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 100


def main() -> None:
    strat_low  = get_pricing_strategy("스마트스토어", "lowest_price")
    strat_norm = get_pricing_strategy("스마트스토어", "normal_sale")
    strat_ad   = get_pricing_strategy("스마트스토어", "cpc_ad")

    wb = openpyxl.Workbook()

    # ── Sheet 1: 최저가전략 ──
    ws1 = wb.active
    ws1.title = "최저가전략"
    _setup_headers(ws1, "lowest_price")
    _fill_lowest_price(ws1, strat_low)

    # ── Sheet 2: 일반판매전략 ──
    ws2 = wb.create_sheet("일반판매전략")
    _setup_headers(ws2, "normal_sale")
    _fill_v_strategy(ws2, strat_norm)

    # ── Sheet 3: 광고전략 ──
    ws3 = wb.create_sheet("광고전략")
    _setup_headers(ws3, "cpc_ad")
    _fill_v_strategy(ws3, strat_ad)

    # ── 정책 시트 ──
    _add_policy_sheet(wb, [
        ("lowest_price", strat_low),
        ("normal_sale", strat_norm),
        ("cpc_ad", strat_ad),
    ])

    out = "DB_save/스마트스토어_3전략_시뮬_v3.xlsx"
    wb.save(out)
    print(f"Saved: {out}")
    print(f"Products: {len(PRODUCTS)}개")

    # ── 콘솔 요약 ──
    rh = strat_low["round_h"]
    print(f"\n=== 최저가전략 (absolute_s, round_h={rh}) ===")
    print(f"{'code':<10}{'cost':>8}{'target_S':>10}{'H_raw':>8}{'H':>7}"
          f"{'I':>9}{'L':>9}{'S':>9}{'V%':>7}  status")
    print("-" * 85)
    for prod in PRODUCTS:
        cost_a = prod["cost_a"]
        ts = _resolve_target_s(cost_a, strat_low["bands"])
        sol = solve_h(cost_a=cost_a, margin_c=strat_low["margin_c"], policy=POLICY,
                      target_s=ts, min_h=strat_low["min_h"], max_h=strat_low["max_h"],
                      round_h=rh)
        fwd = sol["forward"] or {}
        h_disp = f"{sol['h'] or 0:.{rh}f}"
        print(f"{prod['code']:<10}{cost_a:>8,.0f}{ts:>10,.0f}"
              f"{sol['h_raw'] or 0:>8.4f}{h_disp:>7}"
              f"{fwd.get('I',0):>9,.0f}{fwd.get('L',0):>9,.0f}"
              f"{fwd.get('S',0):>9,.0f}{fwd.get('revenue_ratio',0):>6.1f}%"
              f"  {sol['status']}")


if __name__ == "__main__":
    main()
