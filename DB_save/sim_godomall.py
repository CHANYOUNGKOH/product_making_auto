"""고도몰 3-전략 시뮬레이션 엑셀 생성.

실행:
    cd .worktrees/feat-pricing
    python DB_save/sim_godomall.py
"""
from __future__ import annotations

import sys
import os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from DB_save.price_engine import solve_h, solve_h_by_v, _resolve_target_s
from DB_save.pricing_strategies import get_pricing_strategy

# ─── 고도몰 기본 정책 ────────────────────────────────────────────────────────
BASE_POLICY = {
    "commission_rate": 4.0,
    "commission_base": "post_discount",
    "coupon_amount": 50,
    "reward_rate": 0,
    # discount_rate: 원가 구간별로 resolve해서 주입
}

# ─── 샘플 상품 (스마트스토어와 동일 28개) ────────────────────────────────────
COSTS = [
    50, 100, 150, 200, 300, 400, 500,
    700, 1000, 1500, 2000, 3000, 4000, 5000,
    7000, 8000, 10000, 15000, 20000, 30000,
    50000, 70000, 100000, 150000, 200000, 300000, 400000, 500000,
]

PRODUCTS = [
    {"code": f"TEST{i+1:03d}", "cost_a": float(c), "shipping": 0.0}
    for i, c in enumerate(COSTS)
]

# ─── 스타일 ─────────────────────────────────────────────────────────────────
HDR_FONT   = Font(bold=True, size=10, color="000000")
HDR_ORANGE = PatternFill(start_color="F4B942", end_color="F4B942", fill_type="solid")
HDR_BLUE   = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HDR_GREEN  = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
CLAMP_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
THIN = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

BASE_COLS = [
    ("상품코드",        14),
    ("A 원가",         10),
    ("C 마진배수",      10),
    ("E 쿠폰",          8),
    ("G 원가×C",       12),
    ("J 할인쿠폰율",    10),
    ("H 판매가율",      10),
    ("I 마켓등록가",    14),
    ("L 할인후가격",    14),
    ("M 쿠폰정산",      10),
    ("N 마켓수수료",    14),
    ("O 입금액",        12),
    ("S 영업이익",      12),
    ("target",          10),
    ("V 매출대비%",     10),
    ("마켓등록가(10원)", 16),
    ("상태",            12),
]

FMT_INT = "#,##0"
FMT_2DP = "0.00"
FMT_PCT = "0.0%"


def _hdr_fill(strategy_id: str) -> PatternFill:
    return {
        "lowest_price": HDR_ORANGE,
        "normal_sale":  HDR_BLUE,
        "cpc_ad":       HDR_GREEN,
    }[strategy_id]


def _setup_headers(ws, strategy_id: str) -> None:
    fill = _hdr_fill(strategy_id)
    for col_idx, (name, width) in enumerate(BASE_COLS, 1):
        c = ws.cell(row=1, column=col_idx, value=name)
        c.font = HDR_FONT
        c.fill = fill
        c.border = THIN
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 35


def _resolve_j(cost_a: float, discount_bands: list) -> float:
    """원가 구간에 따라 J(할인쿠폰율 %) 반환."""
    return _resolve_target_s(cost_a, discount_bands)


def _make_policy(cost_a: float, discount_bands: list) -> dict:
    """원가에 맞는 J를 주입한 policy dict 반환."""
    j_pct = _resolve_j(cost_a, discount_bands)
    return {**BASE_POLICY, "discount_rate": j_pct}


def _write_row(ws, row_idx: int, row: list, status: str) -> None:
    for col_idx, val in enumerate(row, 1):
        c = ws.cell(row=row_idx, column=col_idx, value=val)
        c.border = THIN
        # 숫자 포맷
        if col_idx in (2, 5, 8, 9, 10, 11, 12, 13, 16):
            c.number_format = FMT_INT
        elif col_idx in (3, 7):
            c.number_format = FMT_2DP
        elif col_idx in (6, 14, 15):
            c.number_format = FMT_PCT
        elif col_idx == 4:
            c.number_format = FMT_INT
        # 경고 색상
        if status != "ok" and col_idx in (7, 17):
            c.fill = CLAMP_FILL


def _fill_lowest_price(ws, strategy: dict) -> None:
    """최저가전략: absolute_s metric, target_S=0."""
    margin_c       = strategy["margin_c"]
    min_h          = strategy["min_h"]
    max_h          = strategy["max_h"]
    round_h        = strategy["round_h"]
    round_mode     = strategy.get("round_mode", "nearest")
    bands          = strategy["bands"]
    discount_bands = strategy["discount_bands"]
    E = BASE_POLICY["coupon_amount"]

    for row_idx, prod in enumerate(PRODUCTS, 2):
        cost_a = prod["cost_a"]
        policy = _make_policy(cost_a, discount_bands)
        J_pct  = policy["discount_rate"]

        target_s = _resolve_target_s(cost_a, bands)   # 항상 0
        sol = solve_h(
            cost_a=cost_a, margin_c=margin_c, policy=policy,
            target_s=target_s, shipping_absorbed=0,
            min_h=min_h, max_h=max_h, round_h=round_h, round_mode=round_mode,
        )
        h   = sol["h"] or 0
        fwd = sol["forward"] or {}
        V_pct = fwd.get("revenue_ratio", 0) or 0

        row = [
            prod["code"],
            cost_a,
            margin_c,
            E,
            cost_a * margin_c,
            J_pct / 100,               # J 할인쿠폰율 (소수)
            h,
            fwd.get("I", 0),
            fwd.get("L", 0),
            E,                         # M 쿠폰정산
            fwd.get("N", 0),
            fwd.get("O", 0),
            fwd.get("S", 0),
            target_s,                  # target_S = 0
            V_pct / 100,
            fwd.get("market_price", 0),
            sol["status"],
        ]
        _write_row(ws, row_idx, row, sol["status"])


def _fill_v_strategy(ws, strategy: dict) -> None:
    """일반판매/광고전략: v_percent metric."""
    margin_c       = strategy["margin_c"]
    min_h          = strategy["min_h"]
    max_h          = strategy["max_h"]
    round_h        = strategy["round_h"]
    round_mode     = strategy.get("round_mode", "nearest")
    bands          = strategy["bands"]
    discount_bands = strategy["discount_bands"]
    E = BASE_POLICY["coupon_amount"]

    for row_idx, prod in enumerate(PRODUCTS, 2):
        cost_a = prod["cost_a"]
        policy = _make_policy(cost_a, discount_bands)
        J_pct  = policy["discount_rate"]

        target_v = _resolve_target_s(cost_a, bands)
        sol = solve_h_by_v(
            cost_a=cost_a, margin_c=margin_c, policy=policy,
            target_v_pct=target_v, shipping_absorbed=0,
            min_h=min_h, max_h=max_h, round_h=round_h, round_mode=round_mode,
        )
        h   = sol["h"] or 0
        fwd = sol["forward"] or {}
        V_pct = fwd.get("revenue_ratio", 0) or 0

        row = [
            prod["code"],
            cost_a,
            margin_c,
            E,
            cost_a * margin_c,
            J_pct / 100,
            h,
            fwd.get("I", 0),
            fwd.get("L", 0),
            E,
            fwd.get("N", 0),
            fwd.get("O", 0),
            fwd.get("S", 0),
            target_v / 100,            # target_V% (소수)
            V_pct / 100,
            fwd.get("market_price", 0),
            sol["status"],
        ]
        _write_row(ws, row_idx, row, sol["status"])


def _add_policy_sheet(wb, strategies: list[tuple[str, dict]]) -> None:
    ws = wb.create_sheet("정책정보")
    rows = [
        ("마켓",       "고도몰"),
        ("수수료율",   f"{BASE_POLICY['commission_rate']}%"),
        ("수수료기준", BASE_POLICY["commission_base"]),
        ("쿠폰 E",    f"{BASE_POLICY['coupon_amount']}원"),
        ("스마일캐시", "0%"),
        ("J 구간",    "원가≤5000:35%, ≤30000:28%, ≤100000:18%, 초과:10%"),
        ("", ""),
    ]
    for sid, strat in strategies:
        rows.append((f"[{sid}] C",     strat["margin_c"]))
        rows.append((f"[{sid}] metric", strat["metric"]))
        rows.append((f"[{sid}] bands", str(strat["bands"])))
        rows.append(("", ""))

    for r, (k, v) in enumerate(rows, 1):
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=str(v))
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 100


def main() -> None:
    strat_low  = get_pricing_strategy("고도몰", "lowest_price")
    strat_norm = get_pricing_strategy("고도몰", "normal_sale")
    strat_ad   = get_pricing_strategy("고도몰", "cpc_ad")

    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "최저가전략"
    _setup_headers(ws1, "lowest_price")
    _fill_lowest_price(ws1, strat_low)

    ws2 = wb.create_sheet("일반판매전략")
    _setup_headers(ws2, "normal_sale")
    _fill_v_strategy(ws2, strat_norm)

    ws3 = wb.create_sheet("광고전략")
    _setup_headers(ws3, "cpc_ad")
    _fill_v_strategy(ws3, strat_ad)

    _add_policy_sheet(wb, [
        ("lowest_price", strat_low),
        ("normal_sale",  strat_norm),
        ("cpc_ad",       strat_ad),
    ])

    out = "DB_save/고도몰_3전략_시뮬_v1.xlsx"
    wb.save(out)
    print(f"Saved: {out}")
    print(f"Products: {len(PRODUCTS)}개")

    # ── 콘솔 요약 (최저가전략) ────────────────────────────────────────────────
    rh = strat_low["round_h"]
    db = strat_low["discount_bands"]
    print(f"\n=== 최저가전략 (absolute_s, S=0, round_h={rh}) ===")
    print(f"{'code':<10}{'cost':>8}{'J%':>6}{'H':>8}"
          f"{'I':>10}{'L':>10}{'S':>8}{'V%':>7}  status")
    print("-" * 75)
    for prod in PRODUCTS:
        cost_a = prod["cost_a"]
        policy = _make_policy(cost_a, db)
        j_pct  = policy["discount_rate"]
        sol = solve_h(
            cost_a=cost_a, margin_c=strat_low["margin_c"], policy=policy,
            target_s=0, min_h=strat_low["min_h"], max_h=strat_low["max_h"],
            round_h=rh, round_mode=strat_low.get("round_mode", "nearest"),
        )
        fwd = sol["forward"] or {}
        h_disp = f"{sol['h'] or 0:.{rh}f}"
        print(f"{prod['code']:<10}{cost_a:>8,.0f}{j_pct:>5.0f}%"
              f"{h_disp:>8}"
              f"{fwd.get('I',0):>10,.0f}{fwd.get('L',0):>10,.0f}"
              f"{fwd.get('S',0):>8,.0f}{fwd.get('revenue_ratio',0):>6.1f}%"
              f"  {sol['status']}")


def _print_boundary_table(strategies: dict) -> None:
    """도매가(원가) 구간 경계점에서 각 전략별 등록가(I) 범위 출력."""
    # J 구간 경계: 원가가 바뀌는 지점 (직전/직후)
    # discount_bands: [5000→35%], [30000→28%], [100000→18%], [초과→10%]
    zones = [
        ("35%", 50,       5_000,   "≤5,000"),
        ("28%", 5_001,   30_000,   "5,001~30,000"),
        ("18%", 30_001, 100_000,   "30,001~100,000"),
        ("10%", 100_001, 500_000,  ">100,000"),
    ]

    strat_low  = strategies["lowest_price"]
    strat_norm = strategies["normal_sale"]
    strat_ad   = strategies["cpc_ad"]

    def get_I(cost_a, strat, metric):
        policy = _make_policy(cost_a, strat["discount_bands"])
        if metric == "absolute_s":
            sol = solve_h(
                cost_a=cost_a, margin_c=strat["margin_c"], policy=policy,
                target_s=0, min_h=strat["min_h"], max_h=strat["max_h"],
                round_h=strat["round_h"], round_mode=strat.get("round_mode", "nearest"),
            )
        else:
            target_v = _resolve_target_s(cost_a, strat["bands"])
            sol = solve_h_by_v(
                cost_a=cost_a, margin_c=strat["margin_c"], policy=policy,
                target_v_pct=target_v, min_h=strat["min_h"], max_h=strat["max_h"],
                round_h=strat["round_h"], round_mode=strat.get("round_mode", "nearest"),
            )
        fwd = sol["forward"] or {}
        return int(fwd.get("market_price", 0))

    print("\n" + "=" * 90)
    print("=== 도매가(원가) 구간별 등록가(I) 범위 - 고도몰 쿠폰 설정 기준 ===")
    print("=" * 90)
    print(f"{'J':>4}  {'원가 구간':>16}  {'등록가 기준':>8}  "
          f"{'최저가전략':>16}  {'일반판매전략':>16}  {'광고전략':>16}")
    print("-" * 90)

    for j_pct, cost_min, cost_max, zone_label in zones:
        for boundary, label in [(cost_min, "최소"), (cost_max, "최대")]:
            i_low  = get_I(boundary, strat_low,  "absolute_s")
            i_norm = get_I(boundary, strat_norm, "v_percent")
            i_ad   = get_I(boundary, strat_ad,   "v_percent")
            print(f"{j_pct:>4}  {zone_label:>16}  {label:>4}(원가{boundary:>7,})"
                  f"  {i_low:>16,}  {i_norm:>16,}  {i_ad:>16,}")
        print()

    print("=" * 90)
    print("※ 고도몰 쿠폰 설정 시 각 전략의 '최대 등록가' 기준으로 구간 상한선 설정 권장")
    print("  (전략 혼용 시 광고전략이 가장 높은 I를 가지므로 광고전략 기준 상한선 사용)")


if __name__ == "__main__":
    main()
    strat_low  = get_pricing_strategy("고도몰", "lowest_price")
    strat_norm = get_pricing_strategy("고도몰", "normal_sale")
    strat_ad   = get_pricing_strategy("고도몰", "cpc_ad")
    _print_boundary_table({
        "lowest_price": strat_low,
        "normal_sale":  strat_norm,
        "cpc_ad":       strat_ad,
    })
