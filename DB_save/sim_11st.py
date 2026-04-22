"""11번가 3-전략 시뮬레이션 엑셀 생성.

실행:
    cd .worktrees/feat-pricing
    python DB_save/sim_11st.py
"""
from __future__ import annotations

import sys
import os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from DB_save.price_engine import solve_h, solve_h_by_v, solve_auto_c, forward_calc, _resolve_target_s
from DB_save.pricing_strategies import get_pricing_strategy

# ─── 11번가 정책 ────────────────────────────────────────────────────────────
POLICY = {
    "commission_rate": 17.5,
    "commission_base": "pre_discount",
    "discount_rate": 50.0,      # J=50% 고정
    "coupon_amount": 120,        # E=120원 고정
    "reward_rate": 0,            # 스마일캐시 없음
}
SHIPPING = 0   # 유배: 구매자 별도 부담

# ─── 샘플 상품 (28개, 원가 50~500,000원) ────────────────────────────────────
COSTS = [
    50, 100, 150, 200, 300, 400, 500,
    700, 1000, 1500, 2000, 3000, 4000, 5000,
    7000, 8000, 10000, 15000, 20000, 30000,
    50000, 70000, 100000, 150000, 200000, 300000, 400000, 500000,
]
PRODUCTS = [
    {"code": f"TEST{i+1:03d}", "cost_a": float(c)}
    for i, c in enumerate(COSTS)
]

# ─── 스타일 ─────────────────────────────────────────────────────────────────
HDR_FONT   = Font(bold=True, size=10, color="FFFFFF")
HDR_ORANGE = PatternFill(start_color="C55A11", end_color="C55A11", fill_type="solid")  # 최저가
HDR_BLUE   = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")  # 일반판매
HDR_GREEN  = PatternFill(start_color="375623", end_color="375623", fill_type="solid")  # 광고
CLAMP_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
THIN = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

# 컬럼 정의
# col: 1=코드 2=원가 3=C 4=J 5=배송 6=G 7=H 8=I 9=L 10=N 11=M(쿠폰) 12=Q(스마일캐시)
#      13=O 14=S 15=target 16=V% 17=S/A 18=L/A 19=등록가10 20=상태
BASE_COLS = [
    ("상품코드",          14),
    ("A 원가",           10),
    ("C 마진배수",        10),
    ("J 할인율",           8),
    ("배송비흡수",        10),
    ("G",                12),
    ("H 판매가율",        10),
    ("I 마켓등록가",      14),
    ("L 할인후가격",      14),
    ("N 마켓수수료",      14),
    ("M 쿠폰",            9),
    ("Q 스마일캐시",      12),
    ("O 입금액",          12),
    ("S 영업이익",        12),
    ("target",            10),
    ("V 매출대비%",       10),
    ("S/A 원가대비이익",  14),
    ("L/A 원가대비매출",  14),
    ("마켓등록가(10원)",  16),
    ("상태",              12),
]

FMT_INT = "#,##0"
FMT_2DP = "0.00"
FMT_PCT = "0.0%"


def _setup_headers(ws, fill: PatternFill) -> None:
    for col_idx, (name, width) in enumerate(BASE_COLS, 1):
        c = ws.cell(row=1, column=col_idx, value=name)
        c.font = HDR_FONT
        c.fill = fill
        c.border = THIN
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 35


def _write_row(ws, row_idx: int, row: list, status: str) -> None:
    for col_idx, val in enumerate(row, 1):
        c = ws.cell(row=row_idx, column=col_idx, value=val)
        c.border = THIN
        if col_idx in (2, 5, 6, 8, 9, 10, 11, 12, 13, 14, 19):
            c.number_format = FMT_INT
        elif col_idx in (3, 7):
            c.number_format = FMT_2DP
        elif col_idx in (4, 15, 16, 17, 18):
            c.number_format = FMT_PCT
        if status not in ("ok",) and col_idx in (7, 20):
            c.fill = CLAMP_FILL


def _fill_lowest(ws, strategy: dict) -> None:
    """최저가전략: absolute_s metric."""
    margin_c   = strategy["margin_c"]
    min_h      = strategy["min_h"]
    max_h      = strategy["max_h"]
    round_h    = strategy["round_h"]
    round_mode = strategy.get("round_mode", "nearest")
    bands      = strategy["bands"]

    row_idx = 2
    prev_I = 0.0
    for prod in PRODUCTS:
        cost_a   = prod["cost_a"]
        target_s = _resolve_target_s(cost_a, bands)
        sol, used_c = solve_auto_c(
            cost_a=cost_a, margin_c=margin_c, policy=POLICY,
            target=target_s, shipping_absorbed=SHIPPING,
            min_h=min_h, max_h=max_h, round_h=round_h, round_mode=round_mode,
            metric="absolute_s",
        )
        h   = sol["h"] or 0
        fwd = sol["forward"] or {}

        # I 단조증가 보장
        cur_I = fwd.get("I", 0) or 0
        if cur_I < prev_I and sol["status"] == "ok":
            G = fwd.get("G", 0) or 1
            h_min_needed = round(prev_I / G + 10 ** (-round_h), round_h)
            h_min_needed = min(h_min_needed, max_h)
            fwd = forward_calc(cost_a, used_c, h_min_needed, POLICY, SHIPPING)
            h = h_min_needed

        V_pct = fwd.get("revenue_ratio", 0) or 0
        L = fwd.get("L", 0) or 0
        S = fwd.get("S", 0) or 0
        prev_I = fwd.get("I", 0) or 0
        sa = (S / cost_a) if cost_a else 0
        la = (L / cost_a) if cost_a else 0

        row = [
            prod["code"], cost_a, used_c,
            POLICY["discount_rate"] / 100,
            SHIPPING,
            fwd.get("G", 0), h,
            fwd.get("I", 0), L,
            fwd.get("N", 0), POLICY["coupon_amount"],
            fwd.get("Q", 0), fwd.get("O", 0), S,
            target_s, V_pct / 100,
            sa, la,
            fwd.get("market_price", 0), sol["status"],
        ]
        _write_row(ws, row_idx, row, sol["status"])
        ws.cell(row=row_idx, column=15).number_format = FMT_INT  # target=원화
        row_idx += 1


def _fill_v(ws, strategy: dict) -> None:
    """일반판매/광고전략: v_percent metric."""
    margin_c = strategy["margin_c"]
    min_h    = strategy["min_h"]
    max_h    = strategy["max_h"]
    round_h  = strategy["round_h"]
    bands    = strategy["bands"]

    row_idx = 2
    prev_I = 0.0
    for prod in PRODUCTS:
        cost_a   = prod["cost_a"]
        target_v = _resolve_target_s(cost_a, bands)
        sol, used_c = solve_auto_c(
            cost_a=cost_a, margin_c=margin_c, policy=POLICY,
            target=target_v, shipping_absorbed=SHIPPING,
            min_h=min_h, max_h=max_h, round_h=round_h,
            metric="v_percent",
        )
        h   = sol["h"] or 0
        fwd = sol["forward"] or {}

        # I 단조증가 보장
        cur_I = fwd.get("I", 0) or 0
        if cur_I < prev_I and sol["status"] == "ok":
            G = fwd.get("G", 0) or 1
            h_min_needed = round(prev_I / G + 10 ** (-round_h), round_h)
            h_min_needed = min(h_min_needed, max_h)
            fwd = forward_calc(cost_a, used_c, h_min_needed, POLICY, SHIPPING)
            h = h_min_needed

        V_pct = fwd.get("revenue_ratio", 0) or 0
        L = fwd.get("L", 0) or 0
        S = fwd.get("S", 0) or 0
        prev_I = fwd.get("I", 0) or 0
        sa = (S / cost_a) if cost_a else 0
        la = (L / cost_a) if cost_a else 0

        row = [
            prod["code"], cost_a, used_c,
            POLICY["discount_rate"] / 100,
            SHIPPING,
            fwd.get("G", 0), h,
            fwd.get("I", 0), L,
            fwd.get("N", 0), POLICY["coupon_amount"],
            fwd.get("Q", 0), fwd.get("O", 0), S,
            target_v / 100, V_pct / 100,
            sa, la,
            fwd.get("market_price", 0), sol["status"],
        ]
        _write_row(ws, row_idx, row, sol["status"])
        row_idx += 1


def _add_policy_sheet(wb, strategies: list) -> None:
    ws = wb.create_sheet("정책정보")
    rows = [
        ("마켓", "11번가"),
        ("수수료율", "17.5%"),
        ("수수료기준", "pre_discount (등록가 기준)"),
        ("할인율 J", "50% 고정"),
        ("쿠폰 E", "120원"),
        ("스마일캐시 P", "0%"),
        ("", ""),
    ]
    for sid, strat in strategies:
        rows.append((f"[{sid}] C", strat["margin_c"]))
        rows.append((f"[{sid}] metric", strat.get("metric", "")))
        rows.append((f"[{sid}] bands", str(strat["bands"])))
        rows.append(("", ""))

    for r, (k, v) in enumerate(rows, 1):
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=str(v))
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100


def main() -> None:
    strat_low  = get_pricing_strategy("11번가", "lowest_price")
    strat_norm = get_pricing_strategy("11번가", "normal_sale")
    strat_ad   = get_pricing_strategy("11번가", "cpc_ad")

    wb = openpyxl.Workbook()

    # Sheet 1: 최저가전략
    ws1 = wb.active
    ws1.title = "최저가전략"
    _setup_headers(ws1, HDR_ORANGE)
    _fill_lowest(ws1, strat_low)

    # Sheet 2: 일반판매전략
    ws2 = wb.create_sheet("일반판매전략")
    _setup_headers(ws2, HDR_BLUE)
    _fill_v(ws2, strat_norm)

    # Sheet 3: 광고전략
    ws3 = wb.create_sheet("광고전략")
    _setup_headers(ws3, HDR_GREEN)
    _fill_v(ws3, strat_ad)

    _add_policy_sheet(wb, [
        ("lowest_price", strat_low),
        ("normal_sale",  strat_norm),
        ("cpc_ad",       strat_ad),
    ])

    out = "DB_save/11번가_3전략_시뮬_v1.xlsx"
    wb.save(out)
    print(f"Saved: {out}")
    print(f"Products: {len(PRODUCTS)}개")

    # ── 콘솔 요약 (최저가전략) ───────────────────────────────────────────────
    rh = strat_low["round_h"]
    print(f"\n=== 11번가 최저가전략 (J=50%, absolute_s, round_h={rh}) ===")
    print(f"{'code':<10}{'cost':>8}{'target_S':>10}{'H_raw':>8}{'H':>6}"
          f"{'I':>10}{'L':>10}{'S':>9}{'V%':>7}  status")
    print("-" * 88)
    for prod in PRODUCTS:
        cost_a   = prod["cost_a"]
        target_s = _resolve_target_s(cost_a, strat_low["bands"])
        sol = solve_h(
            cost_a=cost_a, margin_c=strat_low["margin_c"], policy=POLICY,
            target_s=target_s, shipping_absorbed=SHIPPING,
            min_h=strat_low["min_h"], max_h=strat_low["max_h"],
            round_h=rh,
        )
        fwd = sol["forward"] or {}
        print(f"{prod['code']:<10}{cost_a:>8,.0f}{target_s:>10,.0f}"
              f"{sol['h_raw'] or 0:>8.4f}{sol['h'] or 0:>6.2f}"
              f"{fwd.get('I',0):>10,.0f}{fwd.get('L',0):>10,.0f}"
              f"{fwd.get('S',0):>9,.0f}{fwd.get('revenue_ratio',0):>6.1f}%"
              f"  {sol['status']}")


if __name__ == "__main__":
    main()
