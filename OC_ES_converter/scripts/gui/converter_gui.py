#!/usr/bin/env python3
"""
오너클랜 → 이셀러스 변환기 메인 윈도우
"""

import os
import sys
import subprocess
from pathlib import Path
from typing import Optional
import pandas as pd

from PyQt5.QtWidgets import (
    QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QGridLayout,
    QLabel, QLineEdit, QPushButton, QProgressBar, QTextEdit,
    QFileDialog, QMessageBox, QGroupBox, QFrame, QCheckBox, QScrollArea, QApplication
)
from PyQt5.QtCore import Qt, QThread, pyqtSignal
from PyQt5.QtGui import QFont

# 경로 설정
SCRIPT_DIR = Path(__file__).parent.parent
sys.path.insert(0, str(SCRIPT_DIR))

from settings import Settings

# convert_base.py에서 함수들 import
from convert_base import (
    read_ownerclan,
    convert_ownerclan_to_esellers,
    convert_ownerclan_to_esellers_for_market,
    save_esellers,
    save_esellers_with_validation,
    ESMCategoryMapper,
    CategoryValidator,
    OUTPUT_MARKETS,
    MARKET_CATEGORY_SHEETS,
)
from convert_godomall import (
    convert_ownerclan_to_godomall,
    save_godomall,
    resolve_godomall_template_path,
)
from seo_alt_injector import init_alt_cache, clear_alt_cache


class ConvertWorker(QThread):
    """변환 작업을 수행하는 워커 스레드"""

    progress = pyqtSignal(int)  # 진행률 (0-100)
    status = pyqtSignal(str)    # 상태 메시지
    finished_result = pyqtSignal(dict)  # 결과 데이터
    error = pyqtSignal(str)     # 에러 메시지

    def __init__(self, input_file: str, output_dir: str, category_file: str, selected_markets: list, chunk_size: int = None, seo_alt_enabled: bool = False):
        super().__init__()
        self.input_file = input_file
        self.output_dir = output_dir
        self.category_file = category_file
        self.selected_markets = selected_markets  # 선택된 마켓 ID 목록
        self.chunk_size = chunk_size  # None이면 분할 안함
        self.seo_alt_enabled = seo_alt_enabled

    def run(self):
        try:
            from datetime import datetime

            # 1. 카테고리 매퍼 및 검증기 로드 (5%)
            self.status.emit("카테고리 파일 로드 중...")
            self.progress.emit(5)

            esm_mapper = None
            category_validator = None
            if self.category_file and os.path.exists(self.category_file):
                # 선택된 마켓에 해당하는 시트만 로드 (속도 개선)
                esm_mapper = ESMCategoryMapper(self.category_file, self.selected_markets)
                category_validator = CategoryValidator(self.category_file, self.selected_markets)

            # 2. 입력 파일 읽기 (10%)
            self.status.emit("입력 파일 읽는 중...")
            self.progress.emit(10)

            oc_df = read_ownerclan(self.input_file)
            total_rows = len(oc_df)

            # SEO alt 캐시 초기화 (체크박스 ON일 때만)
            if self.seo_alt_enabled:
                self.status.emit("SEO alt 캐시 로드 중...")
                product_codes = [str(c).strip() for c in oc_df.get("상품코드", []) if str(c).strip()]
                init_alt_cache(product_codes)

            # 출력 디렉토리 생성
            output_path = Path(self.output_dir)
            output_path.mkdir(parents=True, exist_ok=True)

            # 파일명 생성용
            input_name = Path(self.input_file).stem
            date_str = datetime.now().strftime('%Y%m%d')

            # 3. 마켓별 변환 및 저장
            market_results = []
            total_markets = len(self.selected_markets)

            # 통계 집계용
            total_normal = 0
            total_fallback = 0
            total_missing = 0

            for i, market_id in enumerate(self.selected_markets):
                # 마켓 정보 찾기
                market_info = None
                for m in OUTPUT_MARKETS:
                    if m['id'] == market_id:
                        market_info = m
                        break

                if not market_info:
                    continue

                market_name = market_info['name']

                # 진행률 업데이트
                progress_pct = 10 + int(80 * (i / total_markets))
                self.status.emit(f"변환 중... {market_name} ({i+1}/{total_markets})")
                self.progress.emit(progress_pct)

                if market_id == 'godomall':
                    godomall_df, godomall_stats = convert_ownerclan_to_godomall(oc_df)
                    safe_market_name = market_name.replace('/', '_').replace(' ', '_')
                    safe_market_name = safe_market_name.replace('(', '').replace(')', '')
                    godomall_output_path = output_path / 'godomall_output'
                    godomall_output_path.mkdir(parents=True, exist_ok=True)
                    output_file = godomall_output_path / f"godomall_{date_str}_{safe_market_name}_{input_name}.xls"
                    saved_file = save_godomall(godomall_df, resolve_godomall_template_path(), str(output_file))

                    market_result = {
                        'market_id': market_id,
                        'market_name': market_name,
                        'total': len(oc_df),
                        'with_options': 0,
                        'normal_count': godomall_stats['converted_count'],
                        'fallback_count': 0,
                        'missing_count': godomall_stats['skipped_option_count'],
                        'files': [{
                            'type': 'normal',
                            'path': saved_file['path'],
                            'name': Path(saved_file['path']).name,
                            'count': saved_file['count'],
                        }],
                        'skipped_option_products': godomall_stats['skipped_option_products'],
                    }
                    total_normal += godomall_stats['converted_count']
                    total_missing += godomall_stats['skipped_option_count']
                    market_results.append(market_result)
                    continue

                # 마켓별 변환
                basic_df, extended_df = convert_ownerclan_to_esellers_for_market(
                    oc_df, market_id, esm_mapper
                )

                # 파일명 베이스: {날짜}_{마켓명}_{원본파일명}
                safe_market_name = market_name.replace('/', '_').replace(' ', '_')
                safe_market_name = safe_market_name.replace('(', '').replace(')', '')
                base_filename = f"이셀양식기본_{date_str}_{safe_market_name}_{input_name}"

                # 카테고리 검증 대상 마켓인지 확인
                is_validation_target = market_id in MARKET_CATEGORY_SHEETS

                if is_validation_target and category_validator and category_validator.loaded:
                    # 카테고리 검증 후 분리 저장 (chunk_size에 따라 분할)
                    save_result = save_esellers_with_validation(
                        oc_df, basic_df, extended_df,
                        market_id, category_validator,
                        output_path, base_filename,
                        chunk_size=self.chunk_size
                    )

                    # 마켓별 결과 기록
                    market_result = {
                        'market_id': market_id,
                        'market_name': market_name,
                        'total': len(basic_df),
                        'with_options': int((basic_df['선택사항 타입'] != '').sum()),
                        'normal_count': save_result['normal_count'],
                        'fallback_count': save_result['fallback_count'],
                        'missing_count': save_result['missing_count'],
                        'files': [],
                    }

                    # 정상 파일들 (5000개 분할로 여러 파일 가능)
                    for file_info in save_result.get('normal_files', []):
                        market_result['files'].append({
                            'type': 'normal',
                            'path': file_info['path'],
                            'name': Path(file_info['path']).name,
                            'count': file_info['count'],
                        })

                    # 기본카테고리 파일들 (5000개 분할로 여러 파일 가능)
                    for file_info in save_result.get('fallback_files', []):
                        market_result['files'].append({
                            'type': 'fallback',
                            'path': file_info['path'],
                            'name': Path(file_info['path']).name,
                            'count': file_info['count'],
                        })

                    # 누락 리포트 (등록불가 폴더에 저장됨)
                    if save_result['report_file']:
                        market_result['files'].append({
                            'type': 'report',
                            'path': save_result['report_file'],
                            'name': Path(save_result['report_file']).name,
                            'count': save_result['missing_count'],
                        })

                    # 통계 집계
                    total_normal += save_result['normal_count']
                    total_fallback += save_result['fallback_count']
                    total_missing += save_result['missing_count']

                else:
                    # 검증 대상 아님 - 기존 방식으로 저장 (chunk_size에 따라 분할)
                    output_filename = f"{base_filename}.xlsx"
                    output_file = output_path / output_filename
                    saved_files = save_esellers(basic_df, extended_df, str(output_file), chunk_size=self.chunk_size)

                    market_result = {
                        'market_id': market_id,
                        'market_name': market_name,
                        'total': len(basic_df),
                        'with_options': int((basic_df['선택사항 타입'] != '').sum()),
                        'normal_count': len(basic_df),
                        'fallback_count': 0,
                        'missing_count': 0,
                        'files': [],
                    }

                    # 분할된 파일들 추가 (파일별 개수 포함)
                    for file_info in saved_files:
                        market_result['files'].append({
                            'type': 'normal',
                            'path': file_info['path'],
                            'name': Path(file_info['path']).name,
                            'count': file_info['count'],
                        })

                    total_normal += len(basic_df)

                market_results.append(market_result)

            # 5. 결과 집계 (100%)
            self.progress.emit(100)

            result = {
                'total': total_rows,
                'market_count': len(market_results),
                'market_results': market_results,
                'output_dir': str(output_path),
                # 통계 정보
                'stats': {
                    'total_normal': total_normal,
                    'total_fallback': total_fallback,
                    'total_missing': total_missing,
                },
            }

            # SEO alt 캐시 정리
            if self.seo_alt_enabled:
                clear_alt_cache()

            self.status.emit("완료!")
            self.finished_result.emit(result)

        except Exception as e:
            if self.seo_alt_enabled:
                clear_alt_cache()
            import traceback
            self.error.emit(f"{str(e)}\n\n{traceback.format_exc()}")


class ConverterWindow(QMainWindow):
    """메인 윈도우"""

    def __init__(self):
        super().__init__()
        self.settings = Settings()
        self.worker = None
        self.last_output_dir = None
        self.market_checkboxes = {}  # 마켓 체크박스 딕셔너리

        self.init_ui()
        self.load_settings()

    def init_ui(self):
        """UI 초기화"""
        self.setWindowTitle("OC→ES 변환기")
        self.setMinimumSize(650, 780)

        # 중앙 위젯
        central = QWidget()
        self.setCentralWidget(central)
        main_layout = QVBoxLayout(central)
        main_layout.setSpacing(15)
        main_layout.setContentsMargins(20, 20, 20, 20)

        # === 파일 선택 섹션 ===
        file_group = QGroupBox("파일 선택")
        file_layout = QGridLayout(file_group)
        file_layout.setSpacing(10)

        # 입력 파일
        file_layout.addWidget(QLabel("입력 파일:"), 0, 0)
        self.input_edit = QLineEdit()
        self.input_edit.setPlaceholderText("오너클랜 Excel 파일 (*.xlsx)")
        self.input_edit.setMinimumHeight(25)
        file_layout.addWidget(self.input_edit, 0, 1)
        input_btn = QPushButton("찾아보기")
        input_btn.setMinimumHeight(25)
        input_btn.setMinimumWidth(80)
        input_btn.clicked.connect(self.browse_input)
        file_layout.addWidget(input_btn, 0, 2)

        # 입력 파일 상품 수 표시 라벨
        self.product_count_label = QLabel("")
        self.product_count_label.setStyleSheet("color: #0066cc; font-weight: bold;")
        file_layout.addWidget(self.product_count_label, 1, 1)

        # 출력 폴더
        file_layout.addWidget(QLabel("출력 폴더:"), 2, 0)
        self.output_edit = QLineEdit()
        self.output_edit.setPlaceholderText("저장할 폴더 경로")
        self.output_edit.setMinimumHeight(25)
        file_layout.addWidget(self.output_edit, 2, 1)
        output_btn = QPushButton("찾아보기")
        output_btn.setMinimumHeight(25)
        output_btn.setMinimumWidth(80)
        output_btn.clicked.connect(self.browse_output)
        file_layout.addWidget(output_btn, 2, 2)

        # 카테고리 파일
        file_layout.addWidget(QLabel("카테고리 파일:"), 3, 0)
        self.category_edit = QLineEdit()
        self.category_edit.setPlaceholderText("OpenMarketCategory.xlsx (선택)")
        self.category_edit.setMinimumHeight(25)
        self.category_edit.setToolTip(
            "이셀러스 카테고리 파일 (OpenMarketCategory.xlsx)\n\n"
            "최신화 방법:\n"
            "이셀러스주머니 프로그램 > 원본상품 엑셀일괄등록 > 마켓카테고리 엑셀다운"
        )
        file_layout.addWidget(self.category_edit, 3, 1)
        category_btn = QPushButton("찾아보기")
        category_btn.setMinimumHeight(25)
        category_btn.setMinimumWidth(80)
        category_btn.clicked.connect(self.browse_category)
        file_layout.addWidget(category_btn, 3, 2)

        # 카테고리 갱신일 표시
        self.category_date_label = QLabel("")
        self.category_date_label.setStyleSheet("color: #666; font-size: 11px;")
        file_layout.addWidget(self.category_date_label, 4, 1, 1, 2)

        main_layout.addWidget(file_group)

        # === 마켓 선택 섹션 ===
        market_group = QGroupBox("출력 마켓 선택")
        market_layout = QGridLayout(market_group)
        market_layout.setSpacing(8)

        # 마켓 체크박스 생성 (2열 배치)
        for i, market in enumerate(OUTPUT_MARKETS):
            checkbox = QCheckBox(market['name'])
            checkbox.setChecked(market['default'])
            self.market_checkboxes[market['id']] = checkbox

            row = i // 2
            col = i % 2
            market_layout.addWidget(checkbox, row, col)

        # 전체 선택/해제 버튼
        btn_layout = QHBoxLayout()
        select_all_btn = QPushButton("전체 선택")
        select_all_btn.clicked.connect(self.select_all_markets)
        btn_layout.addWidget(select_all_btn)

        deselect_all_btn = QPushButton("전체 해제")
        deselect_all_btn.clicked.connect(self.deselect_all_markets)
        btn_layout.addWidget(deselect_all_btn)

        btn_layout.addStretch()
        market_layout.addLayout(btn_layout, (len(OUTPUT_MARKETS) + 1) // 2, 0, 1, 2)

        main_layout.addWidget(market_group)

        # === 변환 옵션 섹션 ===
        option_group = QGroupBox("변환 옵션")
        option_layout = QHBoxLayout(option_group)

        self.split_checkbox = QCheckBox("5000개 단위로 파일 분할")
        self.split_checkbox.setChecked(False)  # 기본값: 분할 안함
        self.split_checkbox.setToolTip("체크시 5000개 상품마다 별도 파일로 저장\n미체크시 전체 상품을 1개 파일로 저장")
        option_layout.addWidget(self.split_checkbox)

        self.seo_alt_checkbox = QCheckBox("SEO alt 텍스트 자동 삽입")
        self.seo_alt_checkbox.setChecked(self.settings.seo_alt_enabled)
        self.seo_alt_checkbox.setToolTip(
            "체크시 상세설명 HTML 내 이미지 alt 속성에\n"
            "DB의 ST2_JSON 분석 데이터 기반 자연어 설명을 자동 삽입\n"
            "(네이버쇼핑 SEO 보조 시그널 확보)"
        )
        self.seo_alt_checkbox.stateChanged.connect(
            lambda state: setattr(self.settings, 'seo_alt_enabled', bool(state))
        )
        option_layout.addWidget(self.seo_alt_checkbox)
        option_layout.addStretch()

        main_layout.addWidget(option_group)

        # === 메인 실행 버튼 ===
        self.convert_btn = QPushButton("이셀러스양식 변환 실행")
        self.convert_btn.setMinimumHeight(40)
        self.convert_btn.setFont(QFont("", 11, QFont.Bold))
        self.convert_btn.clicked.connect(self.start_convert)
        main_layout.addWidget(self.convert_btn)

        # 오너클랜 원가/배송비 별도 버튼
        self.ownerclan_cost_btn = QPushButton("오너클랜 원가/배송비 엑셀 만들기")
        self.ownerclan_cost_btn.setMinimumHeight(36)
        self.ownerclan_cost_btn.setFont(QFont("", 10, QFont.Bold))
        self.ownerclan_cost_btn.clicked.connect(self.create_ownerclan_cost_shipping_excel)
        main_layout.addWidget(self.ownerclan_cost_btn)

        # === 진행 상황 ===
        progress_group = QGroupBox("진행 상황")
        progress_layout = QVBoxLayout(progress_group)

        self.status_label = QLabel("대기 중")
        progress_layout.addWidget(self.status_label)

        self.progress_bar = QProgressBar()
        self.progress_bar.setValue(0)
        progress_layout.addWidget(self.progress_bar)

        # 시간 정보 라벨
        time_layout = QHBoxLayout()
        self.start_time_label = QLabel("")
        self.start_time_label.setStyleSheet("color: #666;")
        time_layout.addWidget(self.start_time_label)
        time_layout.addStretch()
        self.eta_label = QLabel("")
        self.eta_label.setStyleSheet("color: #0066cc;")
        time_layout.addWidget(self.eta_label)
        progress_layout.addLayout(time_layout)

        main_layout.addWidget(progress_group)

        # === 결과 표시 ===
        result_group = QGroupBox("결과")
        result_layout = QVBoxLayout(result_group)

        self.result_text = QTextEdit()
        self.result_text.setReadOnly(True)
        self.result_text.setMinimumHeight(150)
        result_layout.addWidget(self.result_text)

        # 결과 버튼들
        btn_layout = QHBoxLayout()
        self.open_file_btn = QPushButton("결과 파일 열기")
        self.open_file_btn.clicked.connect(self.open_result_file)
        self.open_file_btn.setEnabled(False)
        btn_layout.addWidget(self.open_file_btn)

        self.open_folder_btn = QPushButton("폴더 열기")
        self.open_folder_btn.clicked.connect(self.open_result_folder)
        self.open_folder_btn.setEnabled(False)
        btn_layout.addWidget(self.open_folder_btn)

        btn_layout.addStretch()
        result_layout.addLayout(btn_layout)

        main_layout.addWidget(result_group)

    def load_settings(self):
        """저장된 설정 로드 (카테고리 파일, 마켓 선택)"""
        # 카테고리 파일 복원
        if self.settings.category_file:
            self.category_edit.setText(self.settings.category_file)
            self.update_category_date_label(self.settings.category_file)

        # 마켓 선택 복원 (저장된 값이 있으면 사용)
        saved_markets = self.settings.selected_markets
        if saved_markets is not None:
            for market_id, checkbox in self.market_checkboxes.items():
                checkbox.setChecked(market_id in saved_markets)

        # 입력/출력 경로는 저장하지 않음 (매번 새로 선택)

    def update_category_date_label(self, category_file_path: str):
        """카테고리 파일의 갱신일 표시 업데이트"""
        if not category_file_path or not os.path.exists(category_file_path):
            self.category_date_label.setText("")
            return

        try:
            import pandas as pd
            # '정보' 시트에서 B3 셀 읽기 (카테고리 갱신일)
            info_df = pd.read_excel(
                category_file_path,
                sheet_name='정보',
                header=None,
                engine='openpyxl'
            )
            # B3 = (2, 1) - 0-indexed
            if len(info_df) > 2 and len(info_df.columns) > 1:
                update_date = info_df.iloc[2, 1]
                if pd.notna(update_date):
                    self.category_date_label.setText(f"카테고리 갱신일: {update_date}")
                    return
        except Exception as e:
            print(f"카테고리 갱신일 읽기 실패: {e}")

        self.category_date_label.setText("")

    def browse_input(self):
        """입력 파일 선택"""
        # 현재 입력된 경로가 있으면 해당 폴더에서 시작
        current_path = self.input_edit.text().strip()
        if current_path and os.path.exists(Path(current_path).parent):
            initial_dir = str(Path(current_path).parent)
        else:
            initial_dir = str(Path.home())

        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "입력 파일 선택",
            initial_dir,
            "Excel Files (*.xlsx);;All Files (*)"
        )
        if file_path:
            self.input_edit.setText(file_path)
            # 경로 저장하지 않음 (매번 새로 선택)

            # 출력 폴더 자동 설정 (입력 파일과 같은 폴더의 output 하위 폴더)
            input_path = Path(file_path)
            output_folder = input_path.parent / '이셀양식변환완료'
            self.output_edit.setText(str(output_folder))

            # 상품 수 읽기
            self.read_product_count(file_path)

    def read_product_count(self, file_path: str):
        """입력 파일에서 상품 수 읽기"""
        try:
            import pandas as pd
            # 오너클랜 파일은 header=1 (2행이 헤더)
            df = pd.read_excel(file_path, engine='openpyxl', header=1)
            count = len(df)
            self.product_count_label.setText(f"상품 수: {count:,}개")
        except Exception as e:
            self.product_count_label.setText("상품 수: 읽기 실패")

    def select_all_markets(self):
        """모든 마켓 선택"""
        for checkbox in self.market_checkboxes.values():
            checkbox.setChecked(True)

    def deselect_all_markets(self):
        """모든 마켓 선택 해제"""
        for checkbox in self.market_checkboxes.values():
            checkbox.setChecked(False)

    def browse_output(self):
        """출력 폴더 선택"""
        # 현재 입력된 경로가 있으면 해당 폴더에서 시작
        current_path = self.output_edit.text().strip()
        if current_path and os.path.exists(current_path):
            initial_dir = current_path
        elif current_path and os.path.exists(Path(current_path).parent):
            initial_dir = str(Path(current_path).parent)
        else:
            initial_dir = str(Path.home())

        folder_path = QFileDialog.getExistingDirectory(
            self,
            "출력 폴더 선택",
            initial_dir
        )
        if folder_path:
            self.output_edit.setText(folder_path)
            # 경로 저장하지 않음 (매번 새로 선택)

    def browse_category(self):
        """카테고리 파일 선택"""
        initial_dir = str(Path(self.category_edit.text()).parent) if self.category_edit.text() else str(Path.home())
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "카테고리 파일 선택",
            initial_dir,
            "Excel Files (*.xlsx);;All Files (*)"
        )
        if file_path:
            self.category_edit.setText(file_path)
            self.settings.category_file = file_path  # 카테고리 파일 경로는 저장
            self.update_category_date_label(file_path)  # 갱신일 업데이트

    def create_ownerclan_cost_shipping_excel(self):
        """입력 파일 기준으로 오너클랜 원가/배송비 엑셀 생성."""
        input_file = self.input_edit.text().strip()
        if not input_file:
            QMessageBox.warning(self, "오류", "입력 파일을 먼저 선택하세요.")
            return
        if not os.path.exists(input_file):
            QMessageBox.warning(self, "오류", f"입력 파일을 찾을 수 없습니다:\n{input_file}")
            return

        self.convert_btn.setEnabled(False)
        self.status_label.setText("오너클랜 원가/배송비 엑셀 생성 중...")
        self.progress_bar.setValue(10)
        self.result_text.clear()
        self.result_text.append(f"입력 파일: {Path(input_file).name}")
        QApplication.processEvents()

        try:
            # 원본 형식: 2행이 헤더, 3행부터 데이터
            df = pd.read_excel(input_file, engine='openpyxl', header=None)
            if df.empty or len(df) < 2:
                raise ValueError("엑셀 헤더(2행)를 찾을 수 없습니다.")

            df.columns = [str(col).strip() if pd.notna(col) else f"Unnamed_{i}" for i, col in enumerate(df.iloc[1])]
            df = df.iloc[2:].reset_index(drop=True)
            self.progress_bar.setValue(30)

            required_columns = ['상품코드', '오너클랜판매가', '배송비', '배송유형', '최대구매수량', '반품배송비']
            available_columns = [str(col).strip() for col in df.columns]

            selected_columns = []
            missing_columns = []
            for req_col in required_columns:
                req_norm = req_col.strip()
                matched_col = None

                for avail in available_columns:
                    if avail == req_norm:
                        matched_col = avail
                        break

                if matched_col is None:
                    req_no_space = req_norm.replace(' ', '').replace('\t', '')
                    for avail in available_columns:
                        avail_no_space = avail.replace(' ', '').replace('\t', '')
                        if avail_no_space == req_no_space:
                            matched_col = avail
                            break

                if matched_col is None:
                    req_low = req_norm.lower().replace(' ', '').replace('\t', '')
                    for avail in available_columns:
                        avail_low = avail.lower().replace(' ', '').replace('\t', '')
                        if req_low in avail_low or avail_low in req_low:
                            matched_col = avail
                            break

                if matched_col is None:
                    missing_columns.append(req_col)
                else:
                    selected_columns.append(matched_col)

            if missing_columns:
                missing_str = ", ".join(missing_columns)
                available_str = ", ".join(available_columns[:20])
                raise ValueError(f"다음 컬럼을 찾을 수 없습니다: {missing_str}\n\n사용 가능한 컬럼(일부): {available_str}")

            result_df = df[selected_columns].copy()
            rename_map = {}
            for i, req_col in enumerate(required_columns):
                if selected_columns[i] != req_col:
                    rename_map[selected_columns[i]] = req_col
            if rename_map:
                result_df = result_df.rename(columns=rename_map)

            self.progress_bar.setValue(55)

            # 계산 컬럼 생성
            result_df['배송비_숫자'] = pd.to_numeric(result_df['배송비'], errors='coerce').fillna(0)
            result_df['오너클랜판매가_숫자'] = pd.to_numeric(result_df['오너클랜판매가'], errors='coerce').fillna(0)

            def calc_3000_price(row):
                shipping = row['배송비_숫자']
                price = row['오너클랜판매가_숫자']
                return price + shipping - 3000 if shipping >= 3000 else price

            def calc_3000_shipping(row):
                shipping = row['배송비_숫자']
                if shipping >= 3000:
                    return 3000
                if shipping > 0:
                    return 3000
                return 0

            def calc_free_price(row):
                shipping = row['배송비_숫자']
                price = row['오너클랜판매가_숫자']
                return price + shipping if shipping > 0 else price

            def calc_free_shipping(row):
                shipping = row['배송비_숫자']
                if shipping > 3000:
                    return shipping
                if shipping > 0:
                    return 3000
                return shipping

            result_df['3000원가'] = result_df.apply(calc_3000_price, axis=1)
            result_df['3000배송비'] = result_df.apply(calc_3000_shipping, axis=1)
            result_df['무배원가'] = result_df.apply(calc_free_price, axis=1)
            result_df['무배배송비'] = result_df.apply(calc_free_shipping, axis=1)

            result_df = result_df.drop(columns=['배송비_숫자', '오너클랜판매가_숫자'])
            base_columns = ['상품코드', '오너클랜판매가', '배송비', '배송유형', '최대구매수량', '반품배송비']
            new_columns = ['3000원가', '3000배송비', '무배원가', '무배배송비']
            result_df = result_df[base_columns + new_columns]
            self.progress_bar.setValue(75)

            output_path = Path(input_file).parent / f"오너클랜원가배송비_{Path(input_file).stem}.xlsx"
            with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
                result_df.to_excel(writer, index=False, sheet_name='Sheet1')

                from openpyxl.styles import PatternFill, Font
                ws = writer.sheets['Sheet1']
                header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
                color1_fill = PatternFill(start_color="4A86E8", end_color="4A86E8", fill_type="solid")
                color2_fill = PatternFill(start_color="6AA84F", end_color="6AA84F", fill_type="solid")
                header_font = Font(bold=True, color="FFFFFF")

                base_col_count = len(base_columns)
                col_3000_price_idx = base_col_count + 1
                col_3000_shipping_idx = base_col_count + 2
                col_free_price_idx = base_col_count + 3
                col_free_shipping_idx = base_col_count + 4

                for col_idx in range(1, len(result_df.columns) + 1):
                    cell = ws.cell(row=1, column=col_idx)
                    cell.font = header_font
                    if col_idx in [col_3000_price_idx, col_3000_shipping_idx]:
                        cell.fill = color1_fill
                    elif col_idx in [col_free_price_idx, col_free_shipping_idx]:
                        cell.fill = color2_fill
                    else:
                        cell.fill = header_fill

                    column_letter = ws.cell(row=1, column=col_idx).column_letter
                    col_name = str(result_df.columns[col_idx - 1])
                    ws.column_dimensions[column_letter].width = min(max(len(col_name) * 1.5, 15), 50)

            self.progress_bar.setValue(100)
            self.status_label.setText("오너클랜 원가/배송비 엑셀 생성 완료")
            self.result_text.append(f"완료: {output_path.name}")
            self.result_text.append(f"행 수: {len(result_df):,}행")
            self.last_output_dir = str(output_path.parent)
            self.open_folder_btn.setEnabled(True)
            self.open_file_btn.setEnabled(False)
            QMessageBox.information(
                self,
                "완료",
                f"오너클랜 원가/배송비 파일 생성 완료\n\n파일명: {output_path.name}\n행 수: {len(result_df):,}개"
            )
        except Exception as e:
            self.status_label.setText("오너클랜 원가/배송비 엑셀 생성 실패")
            self.result_text.append(f"오류: {str(e)}")
            QMessageBox.critical(self, "오류", f"엑셀 생성 실패:\n\n{str(e)}")
        finally:
            self.convert_btn.setEnabled(True)

    def start_convert(self):
        """변환 시작"""
        input_file = self.input_edit.text().strip()
        output_dir = self.output_edit.text().strip()
        category_file = self.category_edit.text().strip()

        # 유효성 검사
        if not input_file:
            QMessageBox.warning(self, "오류", "입력 파일을 선택하세요.")
            return
        if not os.path.exists(input_file):
            QMessageBox.warning(self, "오류", f"입력 파일을 찾을 수 없습니다:\n{input_file}")
            return
        if not output_dir:
            QMessageBox.warning(self, "오류", "출력 폴더 경로를 지정하세요.")
            return

        # 선택된 마켓 확인
        selected_markets = []
        for market_id, checkbox in self.market_checkboxes.items():
            if checkbox.isChecked():
                selected_markets.append(market_id)

        if not selected_markets:
            QMessageBox.warning(self, "오류", "최소 1개 이상의 마켓을 선택하세요.")
            return

        # 선택된 마켓 저장 (다음 실행 시 복원)
        self.settings.selected_markets = selected_markets

        # UI 상태 변경
        self.convert_btn.setEnabled(False)
        self.open_file_btn.setEnabled(False)
        self.open_folder_btn.setEnabled(False)
        self.progress_bar.setValue(0)
        self.result_text.clear()
        self.status_label.setText("변환 준비 중...")

        # 시작시간 기록
        from datetime import datetime
        self.start_time = datetime.now()
        self.start_time_label.setText(f"시작: {self.start_time.strftime('%H:%M:%S')}")
        self.eta_label.setText("")

        # chunk_size 결정 (체크박스 상태에 따라)
        chunk_size = 5000 if self.split_checkbox.isChecked() else None
        seo_alt_enabled = self.seo_alt_checkbox.isChecked()

        # 워커 스레드 시작
        self.worker = ConvertWorker(input_file, output_dir, category_file, selected_markets, chunk_size, seo_alt_enabled)
        self.worker.progress.connect(self.on_progress)
        self.worker.status.connect(self.on_status)
        self.worker.finished_result.connect(self.on_finished)
        self.worker.error.connect(self.on_error)
        self.worker.start()

        self.last_output_dir = output_dir

    def on_progress(self, value: int):
        """진행률 업데이트"""
        self.progress_bar.setValue(value)

        # 예상종료시간 계산
        if value > 5 and hasattr(self, 'start_time'):
            from datetime import datetime, timedelta
            elapsed = (datetime.now() - self.start_time).total_seconds()
            if value > 0:
                total_estimated = elapsed * 100 / value
                remaining = total_estimated - elapsed
                eta = datetime.now() + timedelta(seconds=remaining)
                self.eta_label.setText(f"예상 완료: {eta.strftime('%H:%M:%S')}")

    def on_status(self, message: str):
        """상태 메시지 업데이트"""
        self.status_label.setText(message)

    def on_finished(self, result: dict):
        """변환 완료"""
        self.convert_btn.setEnabled(True)
        self.open_file_btn.setEnabled(False)  # 이제 단일 파일이 아니므로 비활성화
        self.open_folder_btn.setEnabled(True)

        # 완료 시간 표시
        if hasattr(self, 'start_time'):
            from datetime import datetime
            end_time = datetime.now()
            elapsed = (end_time - self.start_time).total_seconds()
            minutes = int(elapsed // 60)
            seconds = int(elapsed % 60)
            self.eta_label.setText(f"완료! (소요시간: {minutes}분 {seconds}초)")

        # 결과 저장 (폴더 열기용)
        self.last_output_dir = result.get('output_dir', self.last_output_dir)

        # 통계 정보
        stats = result.get('stats', {})
        total_normal = stats.get('total_normal', 0)
        total_fallback = stats.get('total_fallback', 0)
        total_missing = stats.get('total_missing', 0)

        # 결과 표시
        lines = [
            f"=== 변환 완료 ===",
            f"입력 상품 수: {result['total']}개",
            f"출력 폴더: {result.get('output_dir', '')}",
            "",
            "=== 카테고리 검증 통계 ===",
            f"  정상 (확장정보 카테고리 사용): {total_normal}개",
            f"  기본카테고리 (esellers 카테고리만 사용): {total_fallback}개",
        ]

        if total_missing > 0:
            lines.append(f"  누락 카테고리 리포트: {total_missing}건")

        lines.append("")
        lines.append("=== 마켓별 생성 파일 ===")

        for mr in result.get('market_results', []):
            lines.append(f"  [{mr['market_name']}]")

            for file_info in mr.get('files', []):
                file_type = file_info['type']
                file_name = file_info['name']
                count = file_info['count']

                if file_type == 'normal':
                    lines.append(f"    ✓ {file_name} ({count}개)")
                elif file_type == 'fallback':
                    lines.append(f"    ⚠ {file_name} ({count}개) - 기본카테고리")
                elif file_type == 'report':
                    lines.append(f"    📋 등록불가/{file_name} ({count}건) - 누락 리포트")

        self.result_text.setText('\n'.join(lines))

        # 통계 알림창 표시 (기본카테고리 또는 누락 있는 경우)
        if total_fallback > 0 or total_missing > 0:
            self.show_stats_alert(result)

    def show_stats_alert(self, result: dict):
        """카테고리 검증 통계 알림창 표시"""
        stats = result.get('stats', {})
        total_normal = stats.get('total_normal', 0)
        total_fallback = stats.get('total_fallback', 0)
        total_missing = stats.get('total_missing', 0)

        # 알림 메시지 구성
        lines = [
            "카테고리 검증 결과",
            "",
            f"✓ 정상: {total_normal}개",
            f"   → 확장정보의 마켓 카테고리 사용",
            "",
        ]

        if total_fallback > 0:
            lines.extend([
                f"⚠ 기본카테고리: {total_fallback}개",
                f"   → 기본정보의 esellers 카테고리만 사용",
                f"   → '_기본카테고리' 파일로 분리됨",
                "",
            ])

        if total_missing > 0:
            lines.extend([
                f"📋 누락 리포트: {total_missing}건",
                f"   → OpenMarketCategory에 미등록된 카테고리",
                f"   → '등록불가' 폴더의 '_누락카테고리' 리포트 참조",
                "",
            ])

        # 마켓별 상세
        lines.append("─" * 30)
        lines.append("마켓별 상세:")

        for mr in result.get('market_results', []):
            normal = mr.get('normal_count', 0)
            fallback = mr.get('fallback_count', 0)
            missing = mr.get('missing_count', 0)

            parts = [f"정상 {normal}개"]
            if fallback > 0:
                parts.append(f"기본카테고리 {fallback}개")
            if missing > 0:
                parts.append(f"누락 {missing}건")

            lines.append(f"  {mr['market_name']}: {' / '.join(parts)}")

        # 알림창 표시
        QMessageBox.information(
            self,
            "카테고리 검증 완료",
            '\n'.join(lines)
        )

    def on_error(self, message: str):
        """에러 발생"""
        self.convert_btn.setEnabled(True)
        self.status_label.setText("오류 발생")
        self.progress_bar.setValue(0)
        QMessageBox.critical(self, "변환 오류", f"변환 중 오류가 발생했습니다:\n\n{message}")

    def open_result_file(self):
        """결과 파일 열기 (현재 비활성화 상태)"""
        pass

    def open_result_folder(self):
        """결과 폴더 열기"""
        if self.last_output_dir and os.path.exists(self.last_output_dir):
            try:
                if sys.platform == 'win32':
                    os.startfile(self.last_output_dir)
                elif sys.platform == 'darwin':
                    subprocess.run(['open', self.last_output_dir])
                else:
                    subprocess.run(['xdg-open', self.last_output_dir])
            except Exception as e:
                QMessageBox.warning(self, "오류", f"폴더를 열 수 없습니다:\n{e}")
